import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
b=pd.read_csv('GSE268904_barcodes_clean.csv',usecols=['sample','condition','sex','age','bmi','endo_specific','endo_ge3','top_lineage','total_umi'])
order={'lean':0,'obese':1,'obese diabetic':2}
nice={'lean':'Lean (no obesity)','obese':'Obese, no diabetes','obese diabetic':'Obese with type 2 diabetes'}
g=b.groupby(['sample','condition','sex','age','bmi'],as_index=False).agg(nuclei=('total_umi','size'),endo=('endo_specific','sum'))
g['pct']=(100*g.endo/g.nuclei).round(1)
g['ord']=g.condition.map(order); g=g.sort_values(['ord','sample'])
by_sample=pd.DataFrame({'Sample':g['sample'],'Group':g.condition.map(nice),'Sex':g.sex.str.capitalize(),'Age (years)':g.age,'BMI':g.bmi,
 'Cells measured':g.nuclei,'Estimated blood-vessel (endothelial) cells':g.endo,'% blood-vessel cells':g.pct})
h=b.groupby('condition').agg(samples=('sample','nunique'),nuclei=('total_umi','size'),endo=('endo_specific','sum')).reset_index()
h['pct']=(100*h.endo/h.nuclei).round(1); h['ord']=h.condition.map(order); h=h.sort_values('ord')
per=g.groupby('condition').pct.agg(['mean','min','max']).round(1)
by_group=pd.DataFrame({'Group':h.condition.map(nice),'People (samples)':h.samples,'Cells measured':h.nuclei,'Estimated blood-vessel cells':h.endo,
 '% blood-vessel cells (all cells pooled)':h.pct,'Average % per person':h.condition.map(per['mean']).values,'Lowest person %':h.condition.map(per['min']).values,'Highest person %':h.condition.map(per['max']).values})
tot=pd.DataFrame({'Group':['All 14 samples'],'People (samples)':[14],'Cells measured':[len(b)],'Estimated blood-vessel cells':[int(b.endo_specific.sum())],
 '% blood-vessel cells (all cells pooled)':[round(100*b.endo_specific.sum()/len(b),1)],'Average % per person':[round(g.pct.mean(),1)],'Lowest person %':[g.pct.min()],'Highest person %':[g.pct.max()]})
by_group=pd.concat([by_group,tot],ignore_index=True)
lab={'endothelial':'Blood-vessel lining (endothelial)','lymphatic_endothelial':'Lymph-vessel lining','adipocyte':'Fat cells (adipocyte)','fibroblast_progenitor':'Connective-tissue / progenitor','immune':'Immune cells','mural':'Vessel wall support (pericyte/muscle)','none':'No marker detected'}
mix=pd.crosstab(b.condition,b.top_lineage,normalize='index').mul(100).round(1).reindex(['lean','obese','obese diabetic']).rename(columns=lab).reset_index()
mix['condition']=mix.condition.map(nice); mix=mix.rename(columns={'condition':'Group'})
lo,hi=g.pct.min(),g.pct.max()
readme=[
['What is this file?'],
['Single-nucleus RNA-seq of human fat (subcutaneous adipose tissue) from 14 people, public dataset GSE268904. About 131,641 individual cell nuclei were measured.'],
[''],
['What do the sheets show?'],
['By sample: one row per person, with how many cells were measured and how many look like blood-vessel lining (endothelial) cells.'],
['By group: the same, added up for each group (lean, obese, obese with type 2 diabetes).'],
['Cell mix (rough): the share of cells whose gene activity looks most like each cell family.'],
[''],
['How was "blood-vessel cell" decided?'],
['A cell is counted when at least 3 of 6 blood-vessel marker genes are active AND blood-vessel genes are its strongest signal, at least twice as strong as any other cell family.'],
[''],
['Important limits'],
['1. These are ESTIMATES from marker genes, not a laboratory cell-type identification. Treat them as approximate.'],
['2. Only 3 people have type 2 diabetes, so comparisons between groups are exploratory and should not be used as proof of a difference.'],
[f'3. People differ a lot (blood-vessel share ranges from {lo}% to {hi}% between samples), so look at the per-person column, not only the group totals.'],
['4. "Cells measured" are individual cell nuclei, not people. The number of independent people is 14.'],
['5. The obese groups are much heavier (BMI about 50) than the lean group (BMI about 23), so group differences may reflect body weight, not diabetes.'],
['6. The Cell mix sheet is very rough: each cell is placed in the family whose marker genes are strongest, using only a few genes per family, so shares can be distorted (for example by RNA floating in the sample). Do not read it as true proportions.']]
with pd.ExcelWriter('GSE268904_summary_for_readers.xlsx',engine='openpyxl') as w:
    pd.DataFrame(readme).to_excel(w,sheet_name='Read me',index=False,header=False)
    by_sample.to_excel(w,sheet_name='By sample',index=False); by_group.to_excel(w,sheet_name='By group',index=False); mix.to_excel(w,sheet_name='Cell mix (rough)',index=False)
wb=load_workbook('GSE268904_summary_for_readers.xlsx'); hdr=PatternFill('solid',fgColor='1F4E79')
for ws in wb:
    if ws.title=='Read me':
        ws.column_dimensions['A'].width=130
        for r in ws.iter_rows():
            c=r[0]; c.alignment=Alignment(wrap_text=True,vertical='top')
            if c.value in ('What is this file?','What do the sheets show?','How was "blood-vessel cell" decided?','Important limits'): c.font=Font(bold=True,size=13,color='1F4E79')
        continue
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    ws.row_dimensions[1].height=48
    for i,col in enumerate(ws.columns,1):
        ws.column_dimensions[get_column_letter(i)].width=max(14,min(34,max(len(str(c.value)) if c.value is not None else 0 for c in col)+2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value,(int,float)) and not isinstance(c.value,bool) and abs(c.value)>=1000: c.number_format='#,##0'
    ws.freeze_panes='A2'
wb.save('GSE268904_summary_for_readers.xlsx')
pd.set_option('display.width',250)
print(by_sample.to_string(index=False)); print(by_group.to_string(index=False)); print(mix.to_string(index=False))
